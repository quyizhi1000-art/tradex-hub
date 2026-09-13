import json,pathlib,csv,collections,hashlib,zipfile
from datetime import datetime
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
p=pathlib.Path('work/state_owned_screen');out=pathlib.Path(r'C:\Users\Admin\Documents\Codex\2026-09-11\new-chat\outputs');out.mkdir(parents=True,exist_ok=True)
d=json.loads((p/'screened.json').read_text(encoding='utf8'));inc=d['included'];exc=d['excluded'];assert len(inc)==1213
straw=json.loads((p/'stock_st_extra.json').read_text(encoding='utf8'));assert len(straw['rows'])==204 and all(x['trade_date']=='20260911' for x in straw['rows']);st={x['ts_code'] for x in straw['rows']}
assert len({x['code'] for x in inc})==len(inc)
assert all(x['code'] not in st and x['master']['market']=='主板' and x['master']['list_status']=='L' and x['code'].endswith(('.SH','.SZ')) and 'ST' not in x['master']['name'].upper() and '退' not in x['master']['name'] for x in inc)
master_old=json.loads((p/'stock_selection_master.json').read_text(encoding='utf8'))['data']['stock_basic'];assert {x['ts_code'] for x in master_old if x['list_status']=='L'}=={x['ts_code'] for x in json.loads((p/'stock_basic_extra.json').read_text(encoding='utf8'))['rows']}
heads=['股票代码','股票名称','完整代码','交易所','板块','地域','行业','分类（按实控人字段）','实控人企业性质原值','实际控制人（来源原值）','入选依据','概念及实控人数据日期','仅昨日概念补入','ST核验日期','资料链接']
def row(x):
 m=x['master'];typ=m.get('act_ent_type');cat=typ if typ in ['中央国企','地方国企'] else '国资参股或相关概念（非控股认定）';ev=x['evidence'];date='；'.join(sorted({z['date'][:10] for z in ev}));prior=all(z['date']=='2026-09-10' for z in ev)
 return [x['code'][:6],m['name'],x['code'],'上交所' if x['code'].endswith('.SH') else '深交所',m['market'],m.get('area'),m.get('industry'),cat,typ if typ is not None else '未提供',m.get('act_name') or '未提供','；'.join(dict.fromkeys(z['origin']+'：'+z['concept'] for z in ev)),date,'是' if prior else '否','2026-09-11',f"https://quote.eastmoney.com/{'sh' if x['code'].endswith('.SH') else 'sz'}{x['code'][:6]}.html"]
rows=[row(x) for x in inc]
wb=Workbook();ov=wb.active;ov.title='筛选口径'
notes=[
('国资相关主板非ST股票筛选','2026-09-11检索'),('总数量',len(inc)),('沪市主板',719),('深市主板',494),('实控人字段：中央国企',381),('实控人字段：地方国企',793),('其余国资参股或相关概念',39),
('纳入规则','满足Tushare实控人企业性质为中央国企/地方国企，或落入已列明的国企改革、国家大基金、证金汇金、国开持股等相关概念之一；合并去重。'),
('证券范围','Tushare当前上市证券（list_status=L），market=主板，且为沪深A股；包括原中小板。'),
('ST排除','2026-09-11 stock_st全量204只，与当前证券简称、当日同花顺成分简称交叉排除；名称含ST、*ST、退市标记亦排除。'),
('本次候选并集',len(inc)+len(exc)),('排除非沪深主板',281),('排除主板ST',45),
('日期说明','同花顺成分数据时间为2026-09-11 10:50；实控人、证券主表、ST名单于2026-09-11获取。东方财富、通达信2026-09-11日批数据尚未返回，以明确标注的2026-09-10已发布快照作补充。'),
('仅昨日概念补入',8),
('完整性边界','已取齐本文件来源表列明的板块成分，并扫描当前上市全市场5563只证券的实控人性质。国资参股采用相关概念标签口径，没有逐公司穿透所有直接、间接股东；未被来源标记的少数参股关系可能遗漏。'),
('分类边界','分类按Tushare实控人性质原值；其余39只以相关概念纳入，不一律认定为国资控股或已确认直接参股。证金、汇金、大基金等财务持股亦在宽口径内。'),
('未扩展的关联','仅有国资云业务合作、仅被国有券商或公募产品持仓、社保养老金持仓，不自动视为国资控股/参股纳入依据。'),
('缺失字段','未提供表示源数据为空；不猜测实际控制人、持股比例或股权链条。'),
('状态时点','非ST是核验日期的状态；停牌未作为额外排除条件。'),
('数据来源','同花顺扶摇（经Tradex SmartRouter）；Tushare（项目配置的限流、校验客户端，含DC/TDX概念日数据）。'),
('官方字段文档','https://tushare.pro/document/1?doc_id=25'),('当日ST文档','https://tushare.pro/document/2?doc_id=397'),('东方财富成分文档','https://tushare.pro/document/2?doc_id=363'),('通达信成分文档','https://tushare.pro/document/2?doc_id=377'),
]
for item in notes:ov.append(item)
ov.column_dimensions['A'].width=30;ov.column_dimensions['B'].width=112;ov.freeze_panes='B2'
for i in range(1,ov.max_row+1):
 ov.cell(i,1).font=Font(bold=True,color='17365D');ov.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top');ov.row_dimensions[i].height=48 if i in [8,9,10,14,16,17,18,19,20,21] else 24
ov.row_dimensions[1].height=32
counter=0
def sheet(name,headers,data,widths=None):
 global counter
 ws=wb.create_sheet(name);ws.append(headers)
 for rr in data:ws.append(rr)
 ws.freeze_panes='C2';ws.sheet_view.showGridLines=False
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=32
 for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=(widths[i-1] if widths and i<=len(widths) else 22)
 if data:
  counter+=1;t=Table(displayName=f'StockTable{counter}',ref=f'A1:{get_column_letter(len(headers))}{ws.max_row}');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
 for rr in ws.iter_rows(min_row=2):
  for c in rr:c.alignment=Alignment(vertical='top',wrap_text=c.column>8)
  rr[0].number_format='@'
 return ws
widths=[12,16,16,12,10,12,18,33,18,55,95,27,18,16,52]
ws=sheet('全部名单_1213只',heads,rows,widths)
for cat,title in [('中央国企','中央国企_381只'),('地方国企','地方国企_793只')]:sheet(title,heads,[r for r in rows if r[7]==cat],widths)
sheet('其他参股及概念_39只',heads,[r for r in rows if r[7] not in ['中央国企','地方国企']],widths)
sheet('昨日概念补充_8只',heads,[r for r in rows if r[12]=='是'],widths)
excludedrows=[]
for x in exc:
 m=x['master'] or {};excludedrows.append([x['code'][:6],m.get('name','未匹配'),x['code'],m.get('market','未匹配'),x['reason'],'；'.join(dict.fromkeys(z['origin']+'：'+z['concept'] for z in x['evidence']))])
sheet('剔除记录_326只',['代码','名称','完整代码','市场','剔除原因','原始入选线索'],excludedrows,[12,16,16,16,30,110])
sources=[]
for suffix in ['TI','DC','TDX']:
 for f in sorted(p.glob('*.'+suffix+'.json')):
  z=json.loads(f.read_text(encoding='utf8'));dt=z['attrs']['provider_as_of'] if suffix=='TI' else '2026-09-10';codes={a['同花顺代码'] if suffix=='TI' else a['con_code'] for a in z['rows']};assert len(codes)==len(z['rows']);sources.append([z['concept'],{'TI':'同花顺扶摇','DC':'Tushare / 东方财富','TDX':'Tushare / 通达信'}[suffix],f.stem,dt,len(z['rows']),len(codes&{x['code'] for x in inc}),'是',hashlib.sha256(f.read_bytes()).hexdigest()])
sources.extend([['全市场实控人性质','Tushare','stock_basic','2026-09-11',5563,1174,'是',hashlib.sha256((p/'stock_basic_extra.json').read_bytes()).hexdigest()],['当日ST清单','Tushare','stock_st','2026-09-11',204,0,'是',hashlib.sha256((p/'stock_st_extra.json').read_bytes()).hexdigest()]])
sheet('来源及数量',['来源主题','提供方','接口或板块','数据日期/时间','原始行数','最终名单命中数（可重叠）','完整获取','原始文件SHA256'],sources,[25,27,24,32,16,29,15,70])
file=out/'国资相关_主板非ST_20260911.xlsx';wb.save(file)
with (out/'国资相关_主板非ST_20260911.csv').open('w',encoding='utf-8-sig',newline='') as f:w=csv.writer(f);w.writerow(heads);w.writerows(rows)
(out/'国资相关_主板非ST_1213只代码.txt').write_text('\n'.join(r[0] for r in rows)+'\n',encoding='utf8')
with zipfile.ZipFile(out/'国资筛选_原始数据与口径_20260911.zip','w',compression=zipfile.ZIP_DEFLATED) as z:
 for f in p.glob('*.json'):z.write(f,arcname=f.name)
 z.writestr('筛选口径.txt','\n\n'.join(str(a)+'：'+str(b) for a,b in notes))
check=load_workbook(file,read_only=True,data_only=False);got=list(check['全部名单_1213只'].iter_rows(min_row=2,values_only=True));assert len(got)==1213 and [r[2] for r in got]==[r[2] for r in rows] and all(isinstance(r[0],str) and len(r[0])==6 for r in got)
assert all('ST' not in r[1].upper() for r in got)
print(json.dumps({'count':len(got),'sheets':check.sheetnames,'files':[str(f) for f in out.glob('国资*')],'checks':'unique codes, main board, listed, current ST exclusion, source completeness, xlsx readback, text codes preserved'},ensure_ascii=False,indent=2))
